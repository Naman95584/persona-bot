"""
triage_raw.py: schema-only inventory of data/raw/ (file size, rows, column names). Never prints values.

    python -m persona_bot.data.triage_raw

CSV: reads the header line and counts newlines in binary chunks (rows can be overcounted if a quoted
field contains a newline, so treat counts as upper bounds). XLSX: openpyxl read_only, header + max_row.
"""
from __future__ import annotations

import csv
import io

from persona_bot.config import RAW_DIR


def csv_info(path) -> dict:
    with open(path, "rb") as fh:
        head = fh.read(1 << 20)
    text = head.decode("utf-8-sig", errors="replace")
    cols = next(csv.reader(io.StringIO(text.split("\n", 1)[0])), [])
    n = 0
    with open(path, "rb") as fh:
        while chunk := fh.read(1 << 24):
            n += chunk.count(b"\n")
    return {"sheet": "-", "rows": max(n - 1, 0), "cols": [c.strip() for c in cols]}


def xlsx_info(path) -> list[dict]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    out = []
    for ws in wb.worksheets:
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        out.append({"sheet": ws.title, "rows": max((ws.max_row or 1) - 1, 0),
                    "cols": [str(c).strip() for c in header if c is not None]})
    wb.close()
    return out


def triage() -> list[dict]:
    rows = []
    for p in sorted(RAW_DIR.rglob("*")):
        if not p.is_file():
            continue
        base = {"file": str(p.relative_to(RAW_DIR)), "mb": round(p.stat().st_size / 1e6, 2)}
        if p.suffix.lower() == ".csv":
            rows.append({**base, **csv_info(p)})
        elif p.suffix.lower() in (".xlsx", ".xlsm"):
            rows += [{**base, **s} for s in xlsx_info(p)]
    return rows


if __name__ == "__main__":
    for r in triage():
        print(f"{r['file']} | sheet={r['sheet']} | {r['mb']} MB | rows~{r['rows']:,} | {len(r['cols'])} cols")
        print("    " + ", ".join(r["cols"]))
